import datetime
import os
import re
from urllib.parse import urlparse

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from cloud_llm import ask_cloud_fallback
from config import RESEARCH_FOLDER, RESEARCH_MAX_RESULTS, TAVILY_API_KEY
from web_search import search_web

# Diferente da conversa normal (respostas curtas, faladas), um relatório de
# pesquisa é pra ser lido depois - por isso usa seu próprio prompt de
# sistema em vez do padrão curto do cloud_llm.
RESEARCH_SYSTEM_PROMPT = (
    "Você escreve relatórios de pesquisa em português do Brasil, não respostas "
    "faladas. Use linguagem clara, impessoal e organizada, com estrutura "
    "acadêmica de leitura inspirada nas normas ABNT. Baseie cada afirmação "
    "somente nas fontes fornecidas; indique citações numéricas como [1] junto "
    "às afirmações correspondentes. Não invente autores, datas, dados ou "
    "referências. Evite Markdown, floreios e repetição."
)


def _build_research_prompt(topic, results):
    snippets = "\n".join(
        f"[{index}] {result.get('title', '')} ({result.get('url', '')}): "
        f"{result.get('content', '')[:700]}"
        for index, result in enumerate(results, start=1)
    )
    return (
        "Escreva o conteúdo de um relatório de pesquisa sobre o tópico. "
        "Organize exatamente nesta sequência: RESUMO (um parágrafo objetivo), "
        "PALAVRAS-CHAVE (3 a 5 termos separados por ponto e vírgula), "
        "1 INTRODUÇÃO, 2 DESENVOLVIMENTO, 3 CONSIDERAÇÕES FINAIS. Use "
        "parágrafos curtos, explique termos técnicos e cite as fontes no texto "
        "com os números fornecidos, por exemplo [1]. Não crie uma seção de "
        "referências: ela será montada pelo programa a partir das URLs reais. "
        "Não use Markdown nem invente dados.\n\n"
        f"Resultados de pesquisa:\n{snippets}\n\n"
        f"Tópico: {topic}\n\n"
        "Conteúdo do relatório:"
    )


def _safe_filename_part(text, max_length=60):
    cleaned = "".join(char for char in text if char.isalnum() or char in " -_").strip()
    return cleaned[:max_length] or "pesquisa"


def _set_run_font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _add_page_number(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def _is_report_heading(line):
    heading = line.strip().strip("#* ").upper()
    if heading in {
        "RESUMO",
        "ABSTRACT",
        "1 INTRODUÇÃO",
        "1. INTRODUÇÃO",
        "2 DESENVOLVIMENTO",
        "2. DESENVOLVIMENTO",
        "3 CONSIDERAÇÕES FINAIS",
        "3. CONSIDERAÇÕES FINAIS",
    }:
        return True
    return bool(re.match(r"^\d+(?:\.\d+)*\s+[A-ZÁÉÍÓÚÂÊÔÃÕÇ]", heading))


def _write_research_docx(topic, report, results, path, researched_at):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3)
    section.left_margin = Cm(3)
    section.bottom_margin = Cm(2)
    section.right_margin = Cm(2)
    _add_page_number(section)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.space_after = Pt(18)
    _set_run_font(title.add_run(topic.strip()), size=14, bold=True)

    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.first_line_indent = Cm(0)
    _set_run_font(metadata.add_run(
        f"Relatório de pesquisa | {researched_at.strftime('%d/%m/%Y')}"
    ), size=10)

    skipped_model_references = False
    for raw_line in report.splitlines():
        line = raw_line.strip().strip("# ")
        if not line:
            continue
        if line.upper().startswith("TÍTULO:"):
            continue
        if line.upper().strip("*:") in {"REFERÊNCIAS", "FONTES"}:
            skipped_model_references = True
            continue
        if skipped_model_references:
            continue
        if line.upper().startswith("PALAVRAS-CHAVE:"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Cm(0)
            label, value = line.split(":", 1)
            _set_run_font(paragraph.add_run(f"{label}:"), bold=True)
            _set_run_font(paragraph.add_run(value))
        elif _is_report_heading(line):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(6)
            _set_run_font(paragraph.add_run(line.strip("#* ")), bold=True)
        elif line.startswith(("- ", "• ")):
            paragraph = document.add_paragraph(line[2:].strip(), style="List Bullet")
            paragraph.paragraph_format.line_spacing = 1.5
        else:
            document.add_paragraph(line)

    references_heading = document.add_paragraph()
    references_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    references_heading.paragraph_format.first_line_indent = Cm(0)
    references_heading.paragraph_format.keep_with_next = True
    references_heading.paragraph_format.space_before = Pt(18)
    _set_run_font(references_heading.add_run("REFERÊNCIAS"), bold=True)

    accessed = researched_at.strftime("%d/%m/%Y")
    for index, result in enumerate(results, start=1):
        url = result.get("url", "")
        domain = urlparse(url).netloc.removeprefix("www.")
        title_text = result.get("title", "Fonte consultada").strip()
        reference = document.add_paragraph()
        reference.alignment = WD_ALIGN_PARAGRAPH.LEFT
        reference.paragraph_format.left_indent = Cm(1.25)
        reference.paragraph_format.first_line_indent = Cm(-1.25)
        reference.paragraph_format.line_spacing = 1.0
        reference.paragraph_format.space_after = Pt(6)
        text = f"[{index}] {title_text}. {domain}. Disponível em: {url}. Acesso em: {accessed}."
        _set_run_font(reference.add_run(text), size=10)

    document.save(path)


def run_research(topic):
    """Pesquisa um tópico na web (Tavily) e usa a nuvem (Groq/Gemini) para
    escrever um relatório .docx com estrutura e apresentação inspiradas na ABNT.
    Retorna o caminho do arquivo salvo. Levanta uma exceção se a busca ou
    a nuvem falharem - quem chama decide a mensagem de erro para o usuário.
    """
    if not TAVILY_API_KEY:
        raise RuntimeError("TAVILY_API_KEY nao configurada - pesquisa na web indisponivel")

    results = search_web(topic, max_results=RESEARCH_MAX_RESULTS)
    if not results:
        raise RuntimeError(f"busca na web nao retornou resultados para '{topic}'")

    prompt = _build_research_prompt(topic, results)
    report = ask_cloud_fallback(prompt, system_prompt=RESEARCH_SYSTEM_PROMPT)

    os.makedirs(RESEARCH_FOLDER, exist_ok=True)
    now = datetime.datetime.now()
    timestamp = now.strftime("%Y-%m-%d %H-%M")
    filename = f"{timestamp} - Pesquisa - {_safe_filename_part(topic)}.docx"
    path = os.path.join(RESEARCH_FOLDER, filename)
    _write_research_docx(topic, report, results, path, now)

    return path
